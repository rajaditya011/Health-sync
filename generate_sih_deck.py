from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Define Colors
NAVY = RGBColor(10, 25, 47)
WHITE = RGBColor(255, 255, 255)
TEAL = RGBColor(0, 150, 136)
LIGHT_TEAL = RGBColor(224, 242, 241)
DARK_GRAY = RGBColor(60, 60, 60)
LIGHT_GRAY = RGBColor(245, 245, 245)
RED = RGBColor(211, 47, 47)
LIGHT_RED = RGBColor(255, 235, 238)
GREEN = RGBColor(46, 125, 50)

def clear_slide(slide):
    # Remove all shapes except the footer if possible, or just remove all editable shapes
    for shape in list(slide.shapes):
        if not (shape.has_text_frame and "Idea Submission" in shape.text):
            slide.shapes._spTree.remove(shape._element)

def add_header(slide, title_text):
    # Dark Navy Header
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.0))
    header.fill.solid()
    header.fill.fore_color.rgb = NAVY
    header.line.fill.background()
    
    # Title Text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.1), Inches(9), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    
    # Accent Line
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.0), Inches(10), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

def add_card(slide, left, top, width, height, title, body, bg_color=LIGHT_GRAY, title_color=TEAL, align=PP_ALIGN.LEFT):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = RGBColor(220, 220, 220)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_top = Inches(0.15)
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = title_color
    p.alignment = align
    
    for line in body:
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.size = Pt(12)
        p2.font.color.rgb = DARK_GRAY
        p2.alignment = align
        if line.startswith("• "):
            p2.level = 1

def add_arrow(slide, left, top, width, height, direction="down"):
    shape_type = MSO_SHAPE.DOWN_ARROW if direction == "down" else MSO_SHAPE.RIGHT_ARROW
    arrow = slide.shapes.add_shape(shape_type, left, top, width, height)
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RGBColor(200, 200, 200)
    arrow.line.fill.background()

def create_presentation():
    prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")
    
    # SLIDE 1: TITLE PAGE
    s1 = prs.slides[0]
    clear_slide(s1)
    
    # Main Title Area
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(3.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    
    tb1 = s1.shapes.add_textbox(Inches(0), Inches(0.5), Inches(10), Inches(0.5))
    p1 = tb1.text_frame.paragraphs[0]
    p1.text = "SMART INDIA HACKATHON 2026"
    p1.font.size = Pt(20)
    p1.font.color.rgb = TEAL
    p1.font.bold = True
    p1.alignment = PP_ALIGN.CENTER
    
    tb2 = s1.shapes.add_textbox(Inches(0), Inches(1.2), Inches(10), Inches(1.0))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = "HEALTHSYNC AI"
    p2.font.size = Pt(54)
    p2.font.color.rgb = WHITE
    p2.font.bold = True
    p2.alignment = PP_ALIGN.CENTER
    
    tb3 = s1.shapes.add_textbox(Inches(0), Inches(2.2), Inches(10), Inches(0.5))
    p3 = tb3.text_frame.paragraphs[0]
    p3.text = "AI-Powered Patient Case-Taking & Intelligent Healthcare Assistance"
    p3.font.size = Pt(18)
    p3.font.color.rgb = WHITE
    p3.alignment = PP_ALIGN.CENTER

    # Info Cards
    add_card(s1, Inches(0.5), Inches(3.8), Inches(4.2), Inches(1.6), "PROJECT DETAILS", 
             ["• Problem ID: SIH26047", "• Title: Patient Case-Taking Software", "• Theme: HEALTH TECH", "• Category: Software"], 
             WHITE, NAVY)
             
    add_card(s1, Inches(5.3), Inches(3.8), Inches(4.2), Inches(1.6), "TEAM: AFTER HOURS (SIH2026-184)", 
             ["Leader: Harsha Kolli (ENG25AM0171)", "Sushanth Hegde | Ashish P | Harsh Kubade", "Aditya Raj | Priyanshu Chauhan"], 
             WHITE, NAVY)
             
    tb4 = s1.shapes.add_textbox(Inches(0), Inches(6.8), Inches(10), Inches(0.5))
    p4 = tb4.text_frame.paragraphs[0]
    p4.text = "School of Engineering, Dayananda Sagar University, Bengaluru - 562112"
    p4.font.size = Pt(14)
    p4.font.color.rgb = DARK_GRAY
    p4.alignment = PP_ALIGN.CENTER
    
    # SLIDE 2: IDEA TITLE
    s2 = prs.slides[1]
    clear_slide(s2)
    add_header(s2, "HEALTHSYNC AI — FROM COMPLEX DATA TO ACTIONABLE CARE")
    
    # Left: Problems
    add_card(s2, Inches(0.5), Inches(1.5), Inches(2.8), Inches(1.0), "01 COMPLEX REPORTS", ["Patients struggle to understand medical reports."], LIGHT_RED, RED)
    add_card(s2, Inches(0.5), Inches(2.7), Inches(2.8), Inches(1.0), "02 RIGHT SPECIALIST", ["Patients may not know which specialist is appropriate."], LIGHT_RED, RED)
    add_card(s2, Inches(0.5), Inches(3.9), Inches(2.8), Inches(1.0), "03 MED. ADHERENCE", ["Patients forget or fail to follow routines."], LIGHT_RED, RED)
    add_card(s2, Inches(0.5), Inches(5.1), Inches(2.8), Inches(1.0), "04 ACCESSIBILITY", ["Language and digital barriers hinder care."], LIGHT_RED, RED)
    
    # Center: Visual Flow
    add_card(s2, Inches(3.8), Inches(1.5), Inches(2.4), Inches(0.8), "PATIENT", [], WHITE, DARK_GRAY, PP_ALIGN.CENTER)
    add_arrow(s2, Inches(4.8), Inches(2.4), Inches(0.4), Inches(0.4))
    add_card(s2, Inches(3.8), Inches(2.9), Inches(2.4), Inches(1.6), "HealthSync AI", ["Report Intelligence", "3D Digital Twin", "Voice AI", "Specialist Matching"], LIGHT_TEAL, TEAL, PP_ALIGN.CENTER)
    add_arrow(s2, Inches(4.8), Inches(4.6), Inches(0.4), Inches(0.4))
    add_card(s2, Inches(3.8), Inches(5.1), Inches(2.4), Inches(0.8), "DOCTOR", [], WHITE, DARK_GRAY, PP_ALIGN.CENTER)
    
    # Right: Innovation
    add_card(s2, Inches(6.7), Inches(1.5), Inches(2.8), Inches(1.2), "3D DIGITAL TWIN", ["Visual symptom & pain localization"], LIGHT_TEAL, NAVY)
    add_card(s2, Inches(6.7), Inches(3.0), Inches(2.8), Inches(1.2), "VOICE-FIRST CARE", ["English/Hindi hands-free interaction"], LIGHT_TEAL, NAVY)
    add_card(s2, Inches(6.7), Inches(4.5), Inches(2.8), Inches(1.2), "AFFORDABLE CARE", ["Jan Aushadhi generic alternatives"], LIGHT_TEAL, NAVY)
    
    # Bottom Banner
    banner = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.5), Inches(10), Inches(0.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    tf_banner = banner.text_frame
    p_banner = tf_banner.paragraphs[0]
    p_banner.text = "AI assists. Doctor decides."
    p_banner.font.size = Pt(20)
    p_banner.font.color.rgb = WHITE
    p_banner.font.bold = True
    p_banner.alignment = PP_ALIGN.CENTER
    
    # SLIDE 3: TECHNICAL APPROACH
    s3 = prs.slides[2]
    clear_slide(s3)
    add_header(s3, "TECHNICAL ARCHITECTURE")
    
    add_card(s3, Inches(1.0), Inches(1.5), Inches(8.0), Inches(0.8), "1. PATIENT INPUT", ["Report Upload  |  Voice Interaction  |  3D Digital Twin"], LIGHT_GRAY, NAVY, PP_ALIGN.CENTER)
    add_arrow(s3, Inches(4.8), Inches(2.4), Inches(0.4), Inches(0.3))
    add_card(s3, Inches(1.0), Inches(2.8), Inches(8.0), Inches(1.2), "2. AI PROCESSING", ["• SambaNova Vision AI → OCR / medical report extraction", "• Groq LLaMA → NLP / contextual analysis / specialist matching"], LIGHT_TEAL, TEAL, PP_ALIGN.CENTER)
    add_arrow(s3, Inches(4.8), Inches(4.1), Inches(0.4), Inches(0.3))
    add_card(s3, Inches(1.0), Inches(4.5), Inches(8.0), Inches(0.8), "3. CLINICAL LAYER (Doctor Dashboard)", ["Review AI suggestions  →  Validate  →  Prescribe"], LIGHT_GRAY, NAVY, PP_ALIGN.CENTER)
    add_arrow(s3, Inches(4.8), Inches(5.4), Inches(0.4), Inches(0.3))
    add_card(s3, Inches(1.0), Inches(5.8), Inches(8.0), Inches(0.8), "4. PATIENT OUTPUT", ["Medication Reminders  |  Diet Guidance  |  Health-Risk Insights  |  Emergency SOS"], LIGHT_GRAY, GREEN, PP_ALIGN.CENTER)
    
    # Tech strip
    ts = s3.shapes.add_textbox(Inches(0), Inches(6.8), Inches(10), Inches(0.5))
    pts = ts.text_frame.paragraphs[0]
    pts.text = "TECH STACK: Next.js 14 • React • Three.js • Framer Motion • Tailwind CSS • Groq API • LLaMA • SambaNova Vision • Web Speech API"
    pts.font.size = Pt(11)
    pts.font.color.rgb = DARK_GRAY
    pts.font.bold = True
    pts.alignment = PP_ALIGN.CENTER
    
    # SLIDE 4: FEASIBILITY AND VIABILITY
    s4 = prs.slides[3]
    clear_slide(s4)
    add_header(s4, "FEASIBILITY, RISKS & MITIGATION")
    
    add_card(s4, Inches(0.5), Inches(1.5), Inches(2.8), Inches(4.5), "WHY FEASIBLE", [
        "MODULAR ARCHITECTURE", "Cloud-based modular components allow independent development and scaling.", "",
        "FAST AI INFERENCE", "Groq and SambaNova APIs support highly responsive AI interactions.", "",
        "ACCESSIBLE DEVICES", "Web-based architecture supports low-end devices effectively."
    ], LIGHT_TEAL, TEAL)
    
    add_card(s4, Inches(3.6), Inches(1.5), Inches(2.8), Inches(4.5), "RISKS", [
        "REPORT QUALITY", "Low-quality / handwritten reports can reduce OCR accuracy.", "",
        "AI HALLUCINATION", "Incorrect AI output could create safety risks.", "",
        "MEDICAL PRIVACY", "Patient data requires strong privacy and consent controls."
    ], LIGHT_RED, RED)
    
    add_card(s4, Inches(6.7), Inches(1.5), Inches(2.8), Inches(4.5), "MITIGATION", [
        "HUMAN-IN-THE-LOOP", "Doctor validates all AI suggestions.", "",
        "ZERO DIAGNOSIS POLICY", "AI strictly assists; it does not independently diagnose.", "",
        "PRIVACY + CONSENT", "Explicit consent and DPDP/ABHA compliance considerations."
    ], LIGHT_GRAY, GREEN)
    
    banner = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(6.5), Inches(10), Inches(0.6))
    banner.fill.solid()
    banner.fill.fore_color.rgb = NAVY
    tf_banner = banner.text_frame
    p_banner = tf_banner.paragraphs[0]
    p_banner.text = "AI + DOCTOR > AI ALONE"
    p_banner.font.size = Pt(22)
    p_banner.font.color.rgb = WHITE
    p_banner.font.bold = True
    p_banner.alignment = PP_ALIGN.CENTER

    # SLIDE 5: IMPLEMENTATION WORKFLOW
    s5 = prs.slides[4]
    clear_slide(s5)
    add_header(s5, "END-TO-END IMPLEMENTATION WORKFLOW")
    
    # Horizontal flow
    add_card(s5, Inches(0.2), Inches(1.5), Inches(1.6), Inches(2.0), "1. INPUT", ["📄 Report", "🎙️ Voice", "🧍 3D Twin"], LIGHT_GRAY, NAVY, PP_ALIGN.CENTER)
    add_arrow(s5, Inches(1.9), Inches(2.3), Inches(0.3), Inches(0.2), "right")
    
    add_card(s5, Inches(2.3), Inches(1.5), Inches(2.0), Inches(2.0), "2. UNDERSTANDING", ["SambaNova Vision", "OCR + extraction", "", "Groq LLaMA", "Context + analysis"], LIGHT_TEAL, TEAL, PP_ALIGN.CENTER)
    add_arrow(s5, Inches(4.4), Inches(2.3), Inches(0.3), Inches(0.2), "right")
    
    add_card(s5, Inches(4.8), Inches(1.5), Inches(2.0), Inches(2.0), "3. CLINICAL AI", ["Risk insights", "Specialist matching", "Recovery planning"], LIGHT_GRAY, NAVY, PP_ALIGN.CENTER)
    add_arrow(s5, Inches(6.9), Inches(2.3), Inches(0.3), Inches(0.2), "right")
    
    add_card(s5, Inches(7.3), Inches(1.5), Inches(2.5), Inches(2.0), "4. VALIDATION", ["Doctor Dashboard", "", "Review", "↓", "Validate", "↓", "Prescribe"], LIGHT_TEAL, NAVY, PP_ALIGN.CENTER)
    
    add_arrow(s5, Inches(8.5), Inches(3.6), Inches(0.2), Inches(0.3), "down")
    add_card(s5, Inches(4.8), Inches(4.0), Inches(5.0), Inches(1.5), "5. PERSONALIZED CARE", ["Medication reminders  |  Diet guidance  |  Recovery tracking"], LIGHT_GRAY, GREEN, PP_ALIGN.CENTER)
    
    # Emergency Branch
    add_card(s5, Inches(0.2), Inches(4.0), Inches(4.2), Inches(1.5), "EMERGENCY SOS", ["Emergency detected → Ambulance support → Family notification"], LIGHT_RED, RED, PP_ALIGN.CENTER)
    
    # Uniqueness Badges
    add_card(s5, Inches(0.5), Inches(6.0), Inches(2.8), Inches(0.6), "", ["★ 3D DIGITAL TWIN"], WHITE, NAVY, PP_ALIGN.CENTER)
    add_card(s5, Inches(3.6), Inches(6.0), Inches(2.8), Inches(0.6), "", ["★ DOCTOR-CONTROLLED AI"], WHITE, NAVY, PP_ALIGN.CENTER)
    add_card(s5, Inches(6.7), Inches(6.0), Inches(2.8), Inches(0.6), "", ["★ DYNAMIC RECOVERY"], WHITE, NAVY, PP_ALIGN.CENTER)

    # SLIDE 6: IMPACT AND BENEFITS
    s6 = prs.slides[5]
    clear_slide(s6)
    add_header(s6, "IMPACT, SCALABILITY & COMPLIANCE")
    
    add_card(s6, Inches(0.5), Inches(1.5), Inches(4.3), Inches(2.2), "SOCIAL IMPACT", [
        "• Better medication adherence.",
        "• Reduced healthcare anxiety.",
        "• More accessible healthcare interaction.",
        "• Support for elderly and chronic patients."
    ], LIGHT_GRAY, TEAL)
    
    add_card(s6, Inches(5.2), Inches(1.5), Inches(4.3), Inches(2.2), "ECONOMIC IMPACT", [
        "• Affordable generic alternatives (Jan Aushadhi).",
        "• Potential reduction in avoidable hospital readmissions.",
        "• Lower friction in healthcare navigation."
    ], LIGHT_GRAY, GREEN)
    
    add_card(s6, Inches(0.5), Inches(4.0), Inches(9.0), Inches(1.8), "FUTURE SCALABILITY PROGRESSION", [
        "HealthSync AI  →  Wearables  →  Real-time Vitals  →  Daily Health Score  →  Preventive Healthcare"
    ], LIGHT_TEAL, NAVY, PP_ALIGN.CENTER)
    
    # Bottom compliance
    tb_comp = s6.shapes.add_textbox(Inches(0), Inches(6.5), Inches(10), Inches(0.5))
    pc = tb_comp.text_frame.paragraphs[0]
    pc.text = "RESEARCH / COMPLIANCE BASIS: MoHFW telemedicine guidance | Jan Aushadhi data | LLaMA / Vision-AI research | ABHA documentation"
    pc.font.size = Pt(10)
    pc.font.color.rgb = DARK_GRAY
    pc.alignment = PP_ALIGN.CENTER
    
    # Remove slides 7 and 8 to keep it exactly 6 slides
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    if len(slides) > 6:
        for i in range(6, len(slides)):
            xml_slides.remove(slides[i])

    prs.save("HealthSync_SIH2026_Pitch_Deck.pptx")
    print("Created HealthSync_SIH2026_Pitch_Deck.pptx")

create_presentation()
