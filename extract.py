from pptx import Presentation
import PyPDF2
import sys

def extract_ppt(file_path):
    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        print(f"--- Slide {i+1} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                print(shape.text)
        print()

def extract_pdf(file_path):
    reader = PyPDF2.PdfReader(file_path)
    for i, page in enumerate(reader.pages):
        print(f"--- Page {i+1} ---")
        print(page.extract_text())
        print()

print("==== PPT ====")
extract_ppt("SIH2026-IDEA-Presentation-Format.pptx")
print("==== PDF ====")
extract_pdf("HealthSync.Ai.pdf")
