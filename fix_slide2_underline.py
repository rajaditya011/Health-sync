from pptx import Presentation
from pptx.util import Inches

prs = Presentation("HealthSync_SIH2026_Perfect_Layout.pptx")
s2 = prs.slides[1]  # Slide 2 (0-indexed)

for shape in s2.shapes:
    # The accent line is 'Rectangle 2' at y=1.300, h=0.040
    if shape.name == "Rectangle 2" and abs(shape.height - Inches(0.04)) < Inches(0.01):
        old_top = shape.top / 914400
        # Title textbox bottom = 0.650 + 0.600 = 1.250
        # Place accent 0.02" below title bottom → y = 1.270
        shape.top = Inches(1.27)
        new_top = shape.top / 914400
        print(f"Accent line moved: y {old_top:.3f} → {new_top:.3f} (gap reduced from 0.050\" to 0.020\")")
        break

prs.save("HealthSync_SIH2026_Perfect_Layout.pptx")
print("Saved.")
