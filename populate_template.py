from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import copy

def replace_text(shape, new_text, font_size=18, bullet=True):
    # clear existing paragraphs
    tf = shape.text_frame
    tf.clear()
    
    for line in new_text.split('\n'):
        if line.strip() == '':
            continue
        p = tf.add_paragraph()
        if line.startswith("- ") or line.startswith("• "):
            p.text = line[2:].strip()
            p.level = 1 if bullet else 0
            p.font.size = Pt(font_size - 2)
        else:
            p.text = line.strip()
            p.level = 0
            p.font.size = Pt(font_size)
            p.font.bold = True
        
        # SIH default text color is often dark grey/black
        p.font.color.rgb = RGBColor(0, 0, 0)

prs = Presentation("SIH2026-IDEA-Presentation-Format.pptx")

# Slide 1: Details
slide1 = prs.slides[0]
for shape in slide1.shapes:
    if shape.name == "TextBox 9":
        text = (
            "Problem Statement ID - SIH26047\n"
            "Problem Statement Title - Patient Case-Taking Software\n"
            "Theme - HEALTH TECH\n"
            "PS Category - Software\n"
            "Team ID - SIH2026-184\n"
            "Team Name - AFTER HOURS\n"
            "Team Leader - Harsha Kolli (ENG25AM0171)\n"
            "Team Members - Sushanth Hegde (ENG25AD0070), Ashish P (ENG25AM0143), "
            "Harsh Kubade (ENG25AM0148), Aditya Raj (ENG25AD0083), Priyanshu Chauhan (ENG25CS0663)"
        )
        tf = shape.text_frame
        tf.clear()
        for line in text.split('\n'):
            p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(16)
            p.font.color.rgb = RGBColor(0, 0, 0)
            p.font.bold = True

# Slide 2: Idea Title
s2_content = (
    "Proposed Solution: HealthSync AI\n"
    "- An integrated platform empowering patients and clinicians with intelligent insights.\n"
    "- Problem Addressed: Complex reports, specialist discovery, and medication non-adherence.\n"
    "- Solution: Scans reports, analyzes conditions, recommends specialists, and tracks recovery.\n"
    "Innovation and Uniqueness:\n"
    "- 3D Digital Twin: Interactive 3D model to visually map symptoms and localized pain.\n"
    "- Multi-lingual Voice Assistant: Hands-free voice (English/Hindi) with Text-To-Speech.\n"
    "- Jan Aushadhi Integration: Recommends affordable generic medicine alternatives."
)

# Slide 3: Tech Approach
s3_content = (
    "Technologies to be Used:\n"
    "- Frontend: Next.js 14, React, Three.js (3D Twin), Framer Motion, Tailwind CSS.\n"
    "- AI Models: Groq API (LLaMA-3 70B for fast NLP), SambaNova API (Llama-3.2 90B Vision).\n"
    "- Voice/Accessibility: Web Speech API for real-time transcription and TTS.\n"
    "- Backend/Architecture: Modular, cloud-based API Routes, Privacy-first design.\n"
    "Methodology and Implementation Process:\n"
    "- Patients interact via the 3D avatar or speak naturally via the Voice Assistant.\n"
    "- SambaNova Vision AI extracts OCR data from uploaded medical reports.\n"
    "- Groq LLaMA analyzes data, assesses urgency, and builds personalized recovery plans.\n"
    "- Integrated ASAP Emergency SOS (Ambulance booking, family notifications)."
)

# Slide 4: Feasibility
s4_content = (
    "Analysis of Feasibility:\n"
    "- High Feasibility: Built on modular, cloud-based architecture optimized for real-world deployment.\n"
    "- Fast Inference: APIs like Groq & SambaNova ensure real-time response on low-end devices.\n"
    "Potential Challenges and Risks:\n"
    "- Extracting medical data from low-quality, handwritten report images.\n"
    "- Preventing AI hallucinations or incorrect medical advice.\n"
    "Strategies for Overcoming Challenges:\n"
    "- Zero Diagnosis Policy: AI is restricted to 'assisting' and never diagnosing without a doctor.\n"
    "- Strict prompt engineering with DPDP Act 2023 & ABHA ethical compliance.\n"
    "- Explicit user consent before recording any online consultations."
)

# Slide 5: Workflow
s5_content = (
    "Architecture & Technology Stack Workflow:\n"
    "- Input Layer: User uploads report / interacts with 3D Twin / speaks to Voice Assistant.\n"
    "- Processing Layer:\n"
    "  - SambaNova Vision extracts OCR data & flags abnormalities.\n"
    "  - Groq LLaMA processes NLP & matches patient to a specialist.\n"
    "  - Doctor Dashboard validates AI suggestions and issues prescriptions.\n"
    "- Output Layer: Medication reminders, personalized diet plans, automated recovery tracking.\n"
    "Highlighting Uniqueness:\n"
    "- Autonomous yet doctor-controlled healthcare.\n"
    "- Dynamic AI Health Risk Score and predictive recovery planning."
)

# Slide 6: Impact
s6_content = (
    "Potential Impact on the Target Audience:\n"
    "- Empowers chronic patients, the elderly, and individuals seeking accessible healthcare.\n"
    "Benefits of the Solution:\n"
    "- Social Impact: Improves medication adherence, reduces anxiety, bridges digital divide.\n"
    "- Economic Impact: Lowers costs via Jan Aushadhi generics; reduces hospital readmissions.\n"
    "Future Scalability:\n"
    "- Wearable device integration (Smartwatches) for real-time vitals monitoring.\n"
    "- AI-based preventive healthcare generating a Daily Health Score."
)

# Slide 7: References
s7_content = (
    "Research & References:\n"
    "- Ministry of Health & Family Welfare guidelines on telemedicine.\n"
    "- Jan Aushadhi scheme data for generic medicine mapping.\n"
    "- LLM integration research based on LLaMA-3 and SambaNova Vision optimization.\n"
    "- ABHA (Ayushman Bharat Health Account) API documentation for data privacy compliance."
)

contents = [s2_content, s3_content, s4_content, s5_content, s6_content, s7_content]

# Update slides 2 to 7
for i in range(1, 7):
    slide = prs.slides[i]
    for shape in slide.shapes:
        if shape.name == "TextBox 8":
            replace_text(shape, contents[i-1])
        # Update team name in Oval
        if shape.name.startswith("Oval") or shape.name.startswith("Your Team"):
            if hasattr(shape, "text"):
                shape.text = "AFTER HOURS"
                for p in shape.text_frame.paragraphs:
                    p.font.color.rgb = RGBColor(255, 255, 255)
                    p.font.size = Pt(14)
                    p.font.bold = True

prs.save("HealthSync_SIH2026_Official_Template.pptx")
print("Saved HealthSync_SIH2026_Official_Template.pptx")
