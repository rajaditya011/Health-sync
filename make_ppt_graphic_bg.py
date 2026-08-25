from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

BG_IMAGE = "/Users/apple/.gemini/antigravity-ide/brain/00986d31-19a6-41b3-95c0-cc258673031f/healthsync_bg_1787590075115.jpg"

def add_background_and_header(slide, is_title=False):
    if os.path.exists(BG_IMAGE):
        bg = slide.shapes.add_picture(BG_IMAGE, Inches(0), Inches(0), Inches(10), Inches(7.5))
        shapes = slide.shapes
        shapes._spTree.remove(bg._element)
        shapes._spTree.insert(2, bg._element)
        
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.2))
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(13, 71, 161)
    header.line.color.rgb = RGBColor(13, 71, 161)
    
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.2), Inches(10), Inches(0.1))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(255, 111, 0)
    accent.line.color.rgb = RGBColor(255, 111, 0)

def format_title(title, text):
    title.text = text
    for p in title.text_frame.paragraphs:
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.font.size = Pt(32)
        p.font.bold = True
        p.alignment = PP_ALIGN.LEFT

def add_title_content_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1] 
    slide = prs.slides.add_slide(slide_layout)
    
    add_background_and_header(slide, is_title=False)
    
    title = slide.shapes.title
    title.left = Inches(0.5)
    title.top = Inches(0.2)
    title.width = Inches(9)
    title.height = Inches(0.8)
    format_title(title, title_text)
    
    bg_content = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.3), Inches(1.5), Inches(9.4), Inches(5.7))
    bg_content.fill.solid()
    bg_content.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_content.line.color.rgb = RGBColor(200, 200, 200)
    
    content = slide.placeholders[1]
    slide.shapes._spTree.remove(content._element)
    slide.shapes._spTree.append(content._element)
    
    content.left = Inches(0.5)
    content.top = Inches(1.7)
    content.width = Inches(9)
    content.height = Inches(5.3)
    tf = content.text_frame
    tf.word_wrap = True
    
    for i, line in enumerate(content_text.split('\n')):
        if line.strip() == '':
            continue
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.text = line.strip().replace("- ", "")
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(50, 50, 50)
        
        if line.startswith("-") or line.startswith("•"):
            p.level = 1
            p.font.size = Pt(16)
        else:
            p.level = 0
            p.font.bold = True
            p.font.color.rgb = RGBColor(0, 137, 123) 

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    
    add_background_and_header(slide, is_title=True)
    
    title = slide.shapes.title
    title.left = Inches(0.5)
    title.top = Inches(0.2)
    title.width = Inches(9)
    title.height = Inches(0.8)
    format_title(title, "SMART INDIA HACKATHON 2026")
    
    bg_sub = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.8), Inches(9), Inches(5.3))
    bg_sub.fill.solid()
    bg_sub.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg_sub.line.color.rgb = RGBColor(200, 200, 200)

    subtitle = slide.placeholders[1]
    slide.shapes._spTree.remove(subtitle._element)
    slide.shapes._spTree.append(subtitle._element)

    subtitle.left = Inches(0.7)
    subtitle.top = Inches(2)
    subtitle.width = Inches(8.6)
    subtitle.height = Inches(5.0)
    
    subtitle_text = (
        "Problem Statement ID: SIH26047\n"
        "Problem Statement Title: Patient Case-Taking Software\n"
        "Theme: HEALTH TECH\n"
        "PS Category: Software\n"
        "Team ID: SIH2026-184\n"
        "Team Name: AFTER HOURS\n"
        "Team Leader: Harsha Kolli (ENG25AM0171)\n"
        "Team Members: Sushanth Hegde (ENG25AD0070), Ashish P (ENG25AM0143), \n"
        "Harsh Kubade (ENG25AM0148), Aditya Raj (ENG25AD0083), Priyanshu Chauhan (ENG25CS0663)\n"
        "School of Engineering, Dayananda Sagar University, Bengaluru - 562112"
    )
    subtitle.text = subtitle_text
    for p in subtitle.text_frame.paragraphs:
        p.font.size = Pt(17)
        p.font.color.rgb = RGBColor(40, 40, 40)
        p.alignment = PP_ALIGN.CENTER
        
    # Slide 2: IDEA TITLE
    s2_content = (
        "IDEA TITLE: HealthSync - AI-Assisted Recovery Journey System\n"
        "Proposed Solution:\n"
        "- HealthSync is an integrated healthcare platform empowering patients and clinicians.\n"
        "- Problem Addressed: Medical reports are complex, finding specialists is hard, and medication non-adherence causes severe risks.\n"
        "- Solution: AI-powered assistant that scans reports, analyzes conditions, recommends specialists, manages online consultations, and tracks recovery.\n"
        "Innovation and Uniqueness:\n"
        "- 3D Digital Twin: An interactive 3D body model for patients to visually map symptoms and localized pain.\n"
        "- Multi-lingual Voice Assistant: Hands-free voice interactions (English/Hindi) with Text-To-Speech for accessibility.\n"
        "- Jan Aushadhi Integration: Recommends affordable generic medicine alternatives."
    )
    add_title_content_slide(prs, "IDEA TITLE & PROPOSED SOLUTION", s2_content)

    # Slide 3: TECHNICAL APPROACH
    s3_content = (
        "Technologies to be Used:\n"
        "- Frontend: Next.js 14, React, Three.js (3D Twin), Framer Motion, Tailwind CSS\n"
        "- AI Models: Groq API (LLaMA-3 70B for fast NLP), SambaNova API (Llama-3.2 90B Vision for OCR/Reports)\n"
        "- Voice/Accessibility: Web Speech API for real-time transcription and Text-to-Speech\n"
        "- Backend/Architecture: Modular, cloud-based API Routes, Privacy-first design\n"
        "Methodology and Implementation Process:\n"
        "- Patients interact via the 3D avatar to report symptoms, or use the voice assistant to speak naturally.\n"
        "- SambaNova Vision AI extracts data from uploaded medical reports and prescriptions.\n"
        "- Groq LLaMA analyzes the data, assesses urgency, and generates a personalized recovery/diet plan.\n"
        "- Integrated ASAP Emergency SOS (Ambulance booking, family notifications, nearest doctor alert)."
    )
    add_title_content_slide(prs, "TECHNICAL APPROACH", s3_content)

    # Slide 4: FEASIBILITY AND VIABILITY
    s4_content = (
        "Analysis of Feasibility:\n"
        "- MVP is highly feasible: Designed as a modular, cloud-based architecture optimized for real-world deployment.\n"
        "- Fast API inference via Groq & SambaNova ensures real-time patient assistance without heavy local computation.\n"
        "Potential Challenges and Risks:\n"
        "- Accurate extraction of medical data from low-quality report images.\n"
        "- Preventing AI hallucinations or incorrect medical advice.\n"
        "Strategies for Overcoming Challenges:\n"
        "- Zero Diagnosis Policy: AI is restricted to 'assisting' and never diagnoses without a doctor’s validation.\n"
        "- Strict prompt engineering with ethical boundaries, DPDP Act 2023 & ABHA compliance.\n"
        "- Explicit user consent before recording consultations."
    )
    add_title_content_slide(prs, "FEASIBILITY AND VIABILITY", s4_content)

    # Slide 5: IMPLEMENTATION WORKFLOW
    s5_content = (
        "Architecture & Technology Stack Workflow:\n"
        "- Input Layer: User uploads report / interacts with 3D Twin / speaks to Voice Assistant.\n"
        "- Processing Layer:\n"
        "   - SambaNova Vision extracts OCR data & flags abnormalities.\n"
        "   - Groq LLaMA processes NLP & matches patient to a specialist.\n"
        "   - Doctor Dashboard validates AI suggestions and issues prescriptions.\n"
        "- Output Layer: Medication reminders, personalized diet & exercise plans, automated recovery tracking.\n"
        "Highlighting Uniqueness:\n"
        "- Autonomous yet doctor-controlled healthcare.\n"
        "- Dynamic AI Health Risk Score and predictive recovery planning.\n"
        "- Seamlessly bridging the digital divide for patients who cannot type (Voice Mode)."
    )
    add_title_content_slide(prs, "IMPLEMENTATION WORKFLOW", s5_content)

    # Slide 6: IMPACT AND BENEFITS
    s6_content = (
        "Potential Impact on the Target Audience:\n"
        "- Targets chronic patients, the elderly, and individuals seeking accessible, stress-free healthcare.\n"
        "Benefits of the Solution:\n"
        "- Social Impact: Improves medication adherence, reduces healthcare anxiety, and provides multilingual support.\n"
        "- Economic Impact: Lowers costs through Jan Aushadhi generic recommendations; reduces hospital readmissions.\n"
        "Future Scalability:\n"
        "- Wearable device integration (Smartwatches) for real-time vitals monitoring (Heart rate, Sleep, SpO2).\n"
        "- AI-based preventive healthcare generating a Daily Health Score to prevent diseases before they become serious."
    )
    add_title_content_slide(prs, "IMPACT AND BENEFITS", s6_content)

    prs.save("HealthSync_SIH2026_Final.pptx")
    print("Presentation saved as HealthSync_SIH2026_Final.pptx")

create_presentation()
