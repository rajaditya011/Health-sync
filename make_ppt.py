from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def add_title_content_slide(prs, title_text, content_text):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.word_wrap = True
    
    for i, line in enumerate(content_text.split('\n')):
        if line.strip() == '':
            continue
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        p.text = line.strip().replace("- ", "")
        p.font.size = Pt(16)
        if line.startswith("-") or line.startswith("•"):
            p.level = 1
        else:
            p.level = 0
            p.font.bold = True

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "SMART INDIA HACKATHON 2026"
    title.text_frame.paragraphs[0].font.size = Pt(36)
    
    subtitle_text = (
        "Problem Statement ID: [Your PS ID]\n"
        "Problem Statement Title: [Your PS Title]\n"
        "Theme: MedTech / Healthcare\n"
        "PS Category: Software\n"
        "Team ID: [Your Team ID]\n"
        "Team Name: [Your Team Name]\n"
        "Team Leader Name: Priyanshu Shaurya (ENG25CS0665)\n"
        "Team Members: Ankit Kumar (ENG25CS0396), Aditya Raj (ENG25AD0083)\n"
        "School of Engineering, Dayananda Sagar University, Bengaluru - 562112."
    )
    subtitle.text = subtitle_text
    for p in subtitle.text_frame.paragraphs:
        p.font.size = Pt(14)
        p.alignment = PP_ALIGN.CENTER
        
    # Slide 2: IDEA TITLE
    s2_content = (
        "IDEA TITLE: HealthSync - AI-Assisted Recovery Journey System\n"
        "Proposed Solution:\n"
        "- HealthSync is an integrated healthcare platform empowering patients and clinicians.\n"
        "- Problem Addressed: Medical reports are complex, finding specialists is hard, and medication non-adherence causes severe risks.\n"
        "- Solution: AI-powered assistant that scans reports, analyzes conditions, recommends specialists, manages online consultations, and tracks recovery.\n"
        "Innovation and Uniqueness:\n"
        "- 3D Digital Twin: An interactive 3D body model for patients to visually map symptoms and localized pain.\n"
        "- Multi-lingual Voice Assistant: Hands-free voice interactions (English/Hindi) with Text-To-Speech for accessibility.\n"
        "- Jan Aushadhi Integration: Recommends affordable generic medicine alternatives."
    )
    add_title_content_slide(prs, "IDEA TITLE & PROPOSED SOLUTION", s2_content)

    # Slide 3: TECHNICAL APPROACH
    s3_content = (
        "Technologies to be Used:\n"
        "- Frontend: Next.js 14, React, Three.js (3D Twin), Framer Motion, Tailwind CSS\n"
        "- AI Models: Groq API (LLaMA-3 70B for fast NLP), SambaNova API (Llama-3.2 90B Vision for OCR/Reports)\n"
        "- Voice/Accessibility: Web Speech API for real-time transcription and Text-to-Speech\n"
        "- Backend/Architecture: Modular, cloud-based API Routes, Privacy-first design\n"
        "Methodology and Implementation Process:\n"
        "- Patients interact via the 3D avatar to report symptoms, or use the voice assistant to speak naturally.\n"
        "- SambaNova Vision AI extracts data from uploaded medical reports and prescriptions.\n"
        "- Groq LLaMA analyzes the data, assesses urgency, and generates a personalized recovery/diet plan.\n"
        "- Integrated ASAP Emergency SOS (Ambulance booking, family notifications, nearest doctor alert)."
    )
    add_title_content_slide(prs, "TECHNICAL APPROACH", s3_content)

    # Slide 4: FEASIBILITY AND VIABILITY
    s4_content = (
        "Analysis of Feasibility:\n"
        "- MVP is highly feasible: Designed as a modular, cloud-based architecture optimized for real-world deployment.\n"
        "- Fast API inference via Groq & SambaNova ensures real-time patient assistance without heavy local computation.\n"
        "Potential Challenges and Risks:\n"
        "- Accurate extraction of medical data from low-quality report images.\n"
        "- Preventing AI hallucinations or incorrect medical advice.\n"
        "Strategies for Overcoming Challenges:\n"
        "- Zero Diagnosis Policy: AI is restricted to 'assisting' and never diagnoses without a doctor’s validation.\n"
        "- Strict prompt engineering with ethical boundaries, DPDP Act 2023 & ABHA compliance.\n"
        "- Explicit user consent before recording consultations."
    )
    add_title_content_slide(prs, "FEASIBILITY AND VIABILITY", s4_content)

    # Slide 5: IMPLEMENTATION WORKFLOW
    s5_content = (
        "Architecture & Technology Stack Workflow:\n"
        "- Input Layer: User uploads report / interacts with 3D Twin / speaks to Voice Assistant.\n"
        "- Processing Layer:\n"
        "   - SambaNova Vision extracts OCR data & flags abnormalities.\n"
        "   - Groq LLaMA processes NLP & matches patient to a specialist.\n"
        "   - Doctor Dashboard validates AI suggestions and issues prescriptions.\n"
        "- Output Layer: Medication reminders, personalized diet & exercise plans, automated recovery tracking.\n"
        "Highlighting Uniqueness:\n"
        "- Autonomous yet doctor-controlled healthcare.\n"
        "- Dynamic AI Health Risk Score and predictive recovery planning.\n"
        "- Seamlessly bridging the digital divide for patients who cannot type (Voice Mode)."
    )
    add_title_content_slide(prs, "IMPLEMENTATION WORKFLOW", s5_content)

    # Slide 6: IMPACT AND BENEFITS
    s6_content = (
        "Potential Impact on the Target Audience:\n"
        "- Targets chronic patients, the elderly, and individuals seeking accessible, stress-free healthcare.\n"
        "Benefits of the Solution:\n"
        "- Social Impact: Improves medication adherence, reduces healthcare anxiety, and provides multilingual support.\n"
        "- Economic Impact: Lowers costs through Jan Aushadhi generic recommendations; reduces hospital readmissions.\n"
        "Future Scalability:\n"
        "- Wearable device integration (Smartwatches) for real-time vitals monitoring (Heart rate, Sleep, SpO2).\n"
        "- AI-based preventive healthcare generating a Daily Health Score to prevent diseases before they become serious."
    )
    add_title_content_slide(prs, "IMPACT AND BENEFITS", s6_content)

    prs.save("HealthSync_SIH2026.pptx")
    print("Presentation saved as HealthSync_SIH2026.pptx")

create_presentation()
