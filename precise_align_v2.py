from pptx import Presentation
from pptx.util import Inches
from PIL import ImageFont
import os

font_path = "/Library/Fonts/Arial.ttf"
if not os.path.exists(font_path):
    font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"

def get_text_width_inches(text, font_size=32):
    if not os.path.exists(font_path):
        return len(text) * 0.25 # rough fallback
    font = ImageFont.truetype(font_path, font_size)
    try:
        bbox = font.getbbox(text)
        width_pixels = bbox[2] - bbox[0]
    except AttributeError:
        width_pixels, _ = font.getsize(text)
    return width_pixels / 72.0

def get_wrapped_width(text, max_width, font_size=32):
    words = text.split(" ")
    lines = []
    current_line = ""
    for word in words:
        test_line = current_line + (" " if current_line else "") + word
        if get_text_width_inches(test_line, font_size) <= max_width:
            current_line = test_line
        else:
            lines.append(current_line)
            current_line = word
    if current_line:
        lines.append(current_line)
        
    max_line_width = max([get_text_width_inches(line, font_size) for line in lines])
    return max_line_width

prs = Presentation("HealthSync_SIH2026_Perfect_Layout.pptx")

for i in range(1, len(prs.slides)):
    slide = prs.slides[i]
    
    title_box = slide.shapes[0]
    underline = slide.shapes[1]
    
    if title_box.has_text_frame:
        title_text = title_box.text_frame.text.replace("\n", " ")
        max_box_width = title_box.width / 914400 # 12.033
        
        # Calculate width of the longest wrapped line
        actual_text_width = get_wrapped_width(title_text, max_box_width, 32)
        
        # Add a tiny bit of padding (0.05 inches) so it doesn't look cut off
        new_line_width = actual_text_width + 0.05
        
        # Ensure it doesn't exceed the box width
        new_line_width = min(new_line_width, max_box_width)
        
        # Set the underline width
        underline.width = Inches(new_line_width)
        print(f"Slide {i+1}: text '{title_text[:20]}...' width set to {new_line_width:.3f} inches")

prs.save("HealthSync_SIH2026_Perfect_Layout.pptx")
print("Precise alignment applied.")
