from pptx import Presentation
from pptx.util import Inches

prs = Presentation("HealthSync_SIH2026_Perfect_Layout.pptx")

for i in range(1, len(prs.slides)):
    slide = prs.slides[i]
    for shape in slide.shapes:
        if shape.name == "Rectangle 2" and abs(shape.height - Inches(0.04)) < Inches(0.01):
            shape.top = Inches(1.27)

prs.save("HealthSync_SIH2026_Perfect_Layout.pptx")
print("All accent lines aligned.")
